"""Build the fictional plant example as native PPTX plus a self-contained PPTD."""
from __future__ import annotations

import argparse
import base64
import json
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt
from template_pack import load_template, inside

ROOT = Path(__file__).resolve().parents[1]
GREEN, CORAL, MUTED, PALE = '#164B3B', '#FF816B', '#52655E', '#EDF7EF'
FONT = 'Reading Text SC'


class Page:
    def __init__(self, slide, kind='content', font_family=FONT):
        self.slide, self.kind, self.elements = slide, kind, []
        self.font_family=font_family

    def record(self, kind, bounds, **kwargs):
        self.elements.append(dict(elementId=f'e{len(self.elements)+1}', elementType=kind,
                                  bounds=list(bounds), **kwargs))

    def shape(self, x, y, w, h, fill='#FFFFFF', circle=False, border=None, rounded=True, radius=None):
        node = self.slide.shapes.add_shape(MSO_SHAPE.OVAL if circle else
                                           (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE),
                                           Pt(x), Pt(y), Pt(w), Pt(h))
        adjustment=.16 if radius is None else min(.5,max(0,radius/min(w,h)))
        if not circle and rounded:
            node.adjustments[0] = adjustment
        node.fill.solid(); node.fill.fore_color.rgb = RGBColor.from_string(fill[1:])
        if border:
            node.line.color.rgb = RGBColor.from_string(border[1:]); node.line.width = Pt(.9)
        else:
            node.line.fill.background()
        kwargs = dict(shapeName='ellipse' if circle else ('roundRect' if rounded else 'rect'), fill=dict(type='solid', color=fill))
        if not circle and rounded:
            kwargs['adjustments'] = [round(adjustment*100000)]
        if border:
            kwargs['border'] = dict(style='solid', width=.9, color=border)
        self.record('shape', [x, y, w, h], **kwargs)

    def text(self, text, x, y, w, h, size=20, color=GREEN, bold=False, center=False, family=None):
        family=family or self.font_family
        box = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        frame = box.text_frame; frame.clear(); frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(text.split('\n')):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
            paragraph.line_spacing = 1.15; paragraph.space_after = Pt(0)
            run = paragraph.add_run(); run.text = line
            run.font.name = family; run.font.size = Pt(size); run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color[1:])
            rpr = run._r.get_or_add_rPr()
            ea = OxmlElement('a:ea'); ea.set('typeface', family); rpr.append(ea)
        self.record('text', [x,y,w,h], content=dict(text=text, fontFamily=family, fontSize=size,
                    bold=bold, color=color, align=['center' if center else 'left','middle'],
                    lineHeight=1.15, wrap=True))

    def line(self, x1, y1, x2, y2, color=GREEN, width=1.2):
        node=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1),Pt(y1),Pt(x2),Pt(y2))
        node.line.color.rgb=RGBColor.from_string(color[1:]);node.line.width=Pt(width)
        w,h=max(abs(x2-x1),.1),max(abs(y2-y1),.1)
        self.record('line',[min(x1,x2),min(y1,y2),w,h],viewBox=[w,h],
                    points=f'{0 if x2>=x1 else w},{0 if y2>=y1 else h} {w if x2>=x1 else 0},{h if y2>=y1 else 0}',
                    border=dict(style='solid',width=width,color=color))

    def image(self, name, x, y, w, h, media, cover=False):
        with Image.open(media/name) as source:
            iw, ih = source.size
        if cover:
            picture = self.slide.shapes.add_picture(str(media/name), Pt(x),Pt(y),width=Pt(w),height=Pt(h))
            if iw/ih > w/h:
                excess = (1 - (w/h)/(iw/ih))/2
                picture.crop_left = picture.crop_right = excess
            else:
                excess = (1 - (iw/ih)/(w/h))/2
                picture.crop_top = picture.crop_bottom = excess
        else:
            scale=min(w/iw,h/ih); rw,rh=iw*scale,ih*scale
            self.slide.shapes.add_picture(str(media/name),Pt(x+(w-rw)/2),Pt(y+(h-rh)/2),width=Pt(rw),height=Pt(rh))
        self.record('image',[x,y,w,h],src='media/'+name,fit=dict(mode='cover' if cover else 'contain'))

    def finish(self):
        transition = OxmlElement('p:transition'); transition.set('spd','med')
        transition.append(OxmlElement('p:fade'))
        self.slide._element.insert_element_before(transition,'p:timing','p:extLst')
        return dict(pageType=self.kind,background=dict(type='solid',color='#FAFCF8'),elements=self.elements)


def write_project(prs, specs, output, basename, font_files, title, body_font=FONT):
    """Write shared native and PPTD representations; font files are relative to output."""
    paths=[]
    for i,spec in enumerate(specs,1):
        rel=f'pages/{i:02}.page';paths.append(rel)
        (output/rel).write_text(yaml.safe_dump(spec,allow_unicode=True,sort_keys=False),encoding='utf-8')
    custom=[]
    for family,relative in font_files.items():
        font_b64=base64.b64encode((output/relative).read_bytes()).decode()
        css=f'@font-face{{font-family:"{family}";font-style:normal;font-weight:400;src:url(data:font/ttf;base64,{font_b64})format("truetype");}}'
        custom.append(dict(family=family,src='data:text/css;base64,'+base64.b64encode(css.encode()).decode()))
    manifest=dict(version='v2',title=title,size=[960,540],customFonts=custom,pages=paths)
    (output/(basename+'.pptd')).write_text(yaml.safe_dump(manifest,allow_unicode=True,sort_keys=False,width=180),encoding='utf-8')
    result=output/(basename+'.pptx');prs.save(result)
    themed=output/'.theme-fonts.pptx'
    with zipfile.ZipFile(result) as source, zipfile.ZipFile(themed,'w',zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            data=source.read(name)
            if name.startswith('ppt/theme/') and name.endswith('.xml'):
                root=ET.fromstring(data)
                for node in root.iter():
                    if node.tag.split('}')[-1] in ['latin','ea','cs','font'] and 'typeface' in node.attrib:
                        node.set('typeface',body_font)
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            elif name=='ppt/presentation.xml':
                root=ET.fromstring(data);root.set('saveSubsetFonts','0')
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            target.writestr(name,data)
    themed.replace(result)
    return result


def build(output: Path, assets: Path | None = None, force: bool = False, template:Path|None=None) -> Path:
    assets = (assets or ROOT/'assets').resolve(); output=output.resolve()
    if output == assets or assets.is_relative_to(output):
        raise ValueError('Output cannot replace or contain the asset directory')
    if output.exists() and any(output.iterdir()) and not force:
        raise ValueError('Output directory is not empty; choose a new directory or use --force')
    template_path=(template or ROOT/'examples/template-extraction/template/template.json').resolve()
    spec=load_template(template_path)
    if output.is_relative_to(template_path.parent) or template_path.parent.is_relative_to(output):
        raise ValueError('Output must be separate from the source template pack')
    if spec['canvas']!=[960,540]:
        raise ValueError('This fixed-content demo expects a 960x540 template')
    if not {'hero-split','process-columns','statement-chapter'}<=set(spec['layouts']):
        raise ValueError('Demo needs hero-split, process-columns and statement-chapter variants')
    for name in ['product-cutout.png']:
        if not (assets/name).is_file():
            raise ValueError(f'Missing example asset: {name}')
    media=output/'media'; pages=output/'pages'
    media.mkdir(parents=True,exist_ok=True); pages.mkdir(parents=True,exist_ok=True)
    shutil.copy2(assets/'product-cutout.png',media/'product-cutout.png')
    shutil.copytree(template_path.parent,output/'template',dirs_exist_ok=True)
    (output/'template/template.json').write_text(json.dumps(spec,ensure_ascii=False,indent=2),encoding='utf8')
    shutil.copy2(inside(template_path.parent,spec['resources']['background']),media/'background.png')
    palette=spec['palette'];types=spec['typography'];components=spec['components']
    primary,accent,body,pale=[palette[k] for k in ['primary','accent','body','pale']]
    font_files={family:'template/'+resource['file'] for family,resource in spec['resources']['fonts'].items()
                if resource.get('file')}
    prs=Presentation();prs.slide_width=Pt(960);prs.slide_height=Pt(540)
    prs.core_properties.author='';prs.core_properties.last_modified_by=''
    prs.core_properties.title='绿芽｜虚构概念演示';prs.core_properties.subject='Reference images to native editable slides'
    specs=[]

    def say(p,text,box,role='body',size=None,color=None,bold=None,center=False):
        style=types[role]
        p.text(text,*box,size if size is not None else style['size'],color or primary,
               style.get('bold',False) if bold is None else bold,center,family=style['family'])
    def card(p,box,fill=None):
        style=components['card']
        p.shape(*box,fill or palette[style['fill']],border=palette[style['border']],radius=style['radius'])
    def footer(p,text):
        cfg=components['footer'];say(p,text,cfg['box'],size=cfg['size'],color=body)
    def relative(box,offset):
        x,y,_,_=box;dx,dy,w,h=offset;return [x+dx,y+dy,w,h]
    def fresh(kind):
        p=Page(prs.slides.add_slide(prs.slide_layouts[6]),kind,font_family=types['body']['family'])
        p.image('background.png',0,0,960,540,media,True)
        tag=components['concept_tag']
        p.shape(*tag['box'],palette[tag['fill']],border=palette['card_border'],radius=tag['radius'])
        say(p,'概念演示',tag['text'],'label',size=14,center=True)
        return p

    p=fresh('cover')
    layout=spec['layouts']['hero-split']['slots']
    card(p,layout['panel'])
    p.image('product-cutout.png',*layout['hero'],media)
    say(p,'绿芽',layout['title'],'display')
    say(p,'桌面植物\n养护提醒系统',layout['subtitle'],'heading',size=types['heading']['size']+2)
    say(p,'看见植物状态，记住每一次照料。',layout['body'],color=body)
    for i,(t,box) in enumerate(zip(['状态查看','养护提醒','记录管理'],layout['features'])):
        p.shape(*box,pale,border=palette['chip_border'],radius=spec['grid']['corner_radius'])
        say(p,t,relative(box,[4,5,box[2]-8,31]),'label',color=accent if i==1 else primary,center=True)
    footer(p,'虚构主题 · 自生成配图 · 用于展示PPT重建方法')
    specs.append(p.finish())

    p=fresh('content')
    layout=spec['layouts']['process-columns']['slots'];cfg=components['process_card']
    say(p,'从植物状态到养护提醒',layout['heading'],'heading')
    p.image('product-cutout.png',*layout['hero'],media)
    steps=[('采集状态','记录植物当前状态'),('整理信息','形成便于查看的资料'),('形成提醒','提示下一次照料'),('记录照料','留下每一次养护记录')]
    for i,(title,detail) in enumerate(steps):
        box=layout['cards'][i];card(p,box)
        say(p,f'{i+1:02}',relative(box,cfg['number']),'numeric',color=palette['soft_primary'],center=True)
        p.shape(*relative(box,cfg['icon']),accent if i==2 else palette['icon_fill'],circle=True)
        say(p,['状态','信息','提醒','记录'][i],relative(box,cfg['icon_text']),'label',color=palette['white'] if i==2 else primary,center=True)
        say(p,title,relative(box,cfg['title']),'heading',size=19,center=True)
        say(p,detail,relative(box,cfg['detail']),size=10.9,color=body,center=True)
        if i<3:say(p,'→',relative(box,cfg['arrow']),'label',size=17,color=palette['arrow'],center=True)
    say(p,'让每一次照料都更清楚。',layout['conclusion'],size=23,color=body,center=True)
    say(p,'流程示意，不包含实测数据',layout['note'],size=11.4,color=body,center=True)
    specs.append(p.finish())

    p=fresh('chapter')
    layout=spec['layouts']['statement-chapter']['slots'];card(p,layout['panel'])
    say(p,'02',layout['number'],'numeric',size=types['numeric']['size']*2,color=palette['soft_primary'])
    say(p,'养护服务\n与交互设计',layout['title'],'heading',size=37)
    say(p,'从状态查看，到提醒与记录',layout['body'],size=18,color=body)
    for i,(t,box) in enumerate(zip(['查看状态','收到提醒','记录照料'],layout['features'])):
        say(p,'• '+t,box,size=13,color=accent if i==1 else primary)
    p.image('product-cutout.png',*layout['hero'],media)
    footer(p,'可复用章节页示例')
    specs.append(p.finish())

    plan={'project':'绿芽｜根据参考图模板制作的示例',
          'style':{'colors':[primary,accent,pale]},
          'design':{'mode':'fidelity','template':{'id':spec['id'],'version':spec['version'],'file':'template/template.json'},
                    'font_profiles':{'from-template':{role:value['family'] for role,value in types.items()}}},
          'slides':[{'id':i,'title':title,'layout':layout,'font_profile':'from-template','status':'concept'}
                    for i,(title,layout) in enumerate([('作品概述','hero-split'),('四步流程','process-columns'),('章节页','statement-chapter')],1)]}
    (output/'design-plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf8')
    return write_project(prs,specs,output,'green-sprout',font_files,'绿芽｜虚构概念演示',types['body']['family'])


def main() -> int:
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,required=True)
    parser.add_argument('--assets-dir',type=Path)
    parser.add_argument('--force',action='store_true')
    parser.add_argument('--template',type=Path,help='Reference-derived template pack; defaults to the bundled example')
    args=parser.parse_args()
    try:
        result=build(args.output_dir,args.assets_dir,args.force,args.template)
    except (OSError,ValueError) as exc:
        parser.error(str(exc))
    print(f'Built 3 editable slides: {result}')
    print('PPTX fonts are declared, not embedded. Install the OFL font or verify substitution before rendering.')
    return 0


if __name__=='__main__':
    raise SystemExit(main())
