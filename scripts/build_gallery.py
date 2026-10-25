"""Build an eight-layout, multi-font design gallery using fictional plant assets."""
from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

from fontTools import subset
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Pt

from build_demo import Page, ROOT, FONT, GREEN, CORAL, MUTED, PALE, write_project

SERIF, PLAY, NUMERIC = 'ZCOOL XiaoWei', 'ZCOOL KuaiLe', 'Barlow Condensed'
PROFILES = {
    'clean': {'heading':FONT,'body':FONT,'numeric':NUMERIC},
    'editorial': {'heading':SERIF,'body':FONT,'numeric':NUMERIC},
    'friendly': {'heading':PLAY,'body':FONT,'numeric':NUMERIC},
}
PAGE_PLAN = [
    ('场景出血与宋体标题','full-bleed','editorial','low','light','作品名称与完整环境主图'),
    ('开放分栏与信息层级','editorial-split','editorial','medium','light','产品主体与三层文字说明'),
    ('沿相同维度对比方案','comparison','clean','medium','mixed','查看、提醒和记录的对应关系'),
    ('四步路径替代等宽卡片','timeline','clean','medium','light','动作的先后关系'),
    ('分层架构说明协作','layered-architecture','clean','high','light','交互、服务与概念采集端的关系'),
    ('大主体与局部引线','annotated-product','editorial','medium','light','主体部件与信息入口'),
    ('童趣字体与深底留白','statement-chapter','friendly','low','dark','一句简短的体验目标'),
    ('错位路线与后续设想','roadmap','editorial','medium','light','分阶段的后续方向'),
]


def build(output:Path, assets:Path|None=None, force:bool=False)->Path:
    assets=(assets or ROOT/'assets').resolve();output=output.resolve()
    if output==assets or assets.is_relative_to(output):
        raise ValueError('Output cannot replace or contain assets')
    if output.exists() and any(output.iterdir()) and not force:
        raise ValueError('Output is not empty; use a new directory or --force')
    media=output/'media';pages=output/'pages';fonts=media/'fonts'
    fonts.mkdir(parents=True,exist_ok=True);pages.mkdir(parents=True,exist_ok=True)
    for name in ['product-cutout.png','lifestyle-scene.png']:
        shutil.copy2(assets/name,media/name)
    prs=Presentation();prs.slide_width=Pt(960);prs.slide_height=Pt(540)
    prs.core_properties.author='';prs.core_properties.last_modified_by=''
    prs.core_properties.title='绿芽｜字体与版式选型图册'
    specs=[]

    def fresh(index,bg='#F8FBF6',dark=False,footer=True):
        p=Page(prs.slides.add_slide(prs.slide_layouts[6]),'chapter' if index==7 else 'content')
        p.shape(0,0,960,540,bg,rounded=False)
        color='#C8DFCF' if dark else MUTED
        if footer:
            p.text(f'{index:02} / {PAGE_PLAN[index-1][0]}',42,492,790,27,11.2,color)
            p.text('概念演示',828,492,90,27,11.2,color,center=True)
        return p

    p=fresh(1,footer=False)
    p.image('lifestyle-scene.png',0,0,960,540,media,True)
    p.text('SPROUT / CONCEPT SYSTEM',48,43,410,34,15,GREEN,family=NUMERIC)
    p.text('绿芽',47,111,366,108,87,GREEN,family=SERIF)
    p.text('让照料，更从容',50,236,386,57,32,GREEN,family=SERIF)
    p.text('看见植物状态，\n记住每一次照料。',51,315,319,76,19,MUTED)
    p.line(51,416,338,416,'#8BAF91',1)
    p.text('状态查看   /   养护提醒   /   记录管理',52,427,371,39,13.7,GREEN)
    # The photo stays in the background; all visible words are native objects.
    p.text('01 / 场景出血与宋体标题',42,492,610,27,11.2,GREEN)
    p.text('概念演示',828,492,90,27,11.2,GREEN,center=True)
    specs.append(p.finish())

    p=fresh(2,'#F2F7EC')
    p.text('FORM / FUNCTION / CARE',44,35,430,32,15,GREEN,family=NUMERIC)
    p.image('product-cutout.png',22,120,421,336,media)
    p.text('从植物状态\n到一次照料',473,66,434,106,38,GREEN,family=SERIF)
    rows=[('01','知道状态','让需要关注的信息集中呈现。'),
          ('02','收到提醒','把下一次照料变成明确提示。'),
          ('03','留下记录','让之前的养护过程便于回看。')]
    for i,(n,t,b) in enumerate(rows):
        y=210+i*84
        p.text(n,474,y,48,39,32,CORAL if i==1 else '#75A07D',family=NUMERIC)
        p.text(t,538,y,355,35,23,GREEN,bold=True)
        p.text(b,538,y+36,355,32,15.5,MUTED)
        if i<2:p.line(537,y+75,902,y+75,'#D7E3D3',.8)
    specs.append(p.finish())

    p=fresh(3)
    p.text('从凭记忆，到有记录',43,37,855,69,35,GREEN,bold=True)
    p.shape(43,132,394,318,PALE,rounded=False)
    p.shape(520,132,397,318,GREEN,rounded=False)
    p.text('情境假设',68,152,341,45,25,GREEN,bold=True)
    p.text('概念方案',548,152,342,45,25,'#FFFFFF',bold=True)
    before=['想起时才查看','不记得上次照料','信息分散在各处']
    after=['状态集中查看','提醒后再处理','记录便于回看']
    for i,(a,b) in enumerate(zip(before,after)):
        y=230+i*64
        p.text(f'{i+1:02}',68,y,49,36,24,'#75A07D',family=NUMERIC)
        p.text(a,126,y,281,39,20,GREEN)
        p.text(f'{i+1:02}',548,y,49,36,24,'#AFCFB8',family=NUMERIC)
        p.text(b,606,y,278,39,20,'#FFFFFF')
    p.text('→',455,269,45,51,30,CORAL,center=True)
    p.text('沿相同维度比较：查看、提醒、记录',47,454,855,32,13.8,MUTED)
    specs.append(p.finish())

    p=fresh(4,'#FFFFFF')
    p.text('用路径，讲清一次照料',44,35,860,68,34,GREEN,bold=True)
    p.text('PROCESS / FOUR STEPS',47,112,710,29,14,'#78A080',family=NUMERIC)
    p.line(118,282,842,282,'#A8CBB1',2.5)
    steps=[('采集状态','记录当前植物状态'),('整理信息','把资料组织成可查看内容'),
           ('形成提醒','提示下一次需要关注的事'),('记录照料','把这次操作留在记录中')]
    for i,(t,b) in enumerate(steps):
        cx=118+i*(724/3);top=i%2==0
        p.shape(cx-24,258,48,48,CORAL if i==2 else GREEN,circle=True)
        p.text(f'{i+1:02}',cx-22,261,44,40,24,'#FFFFFF',center=True,family=NUMERIC)
        y=172 if top else 339
        p.text(t,cx-100,y,200,39,23,GREEN,bold=True,center=True)
        p.text(b,cx-106,y+44,212,47,15.3,MUTED,center=True)
        p.line(cx,245 if top else 319,cx,253 if top else 333,'#A8CBB1',1)
    p.text('流程示意，不包含实测数据',48,447,864,31,13.7,MUTED)
    specs.append(p.finish())

    p=fresh(5,'#F4F8F1')
    p.text('分层表达，让关系更清楚',43,35,866,68,34,GREEN,bold=True)
    layers=[('01','用户交互','查看状态   /   收到提醒','用户可接触的界面与动作'),
            ('02','状态与提醒','信息组织   /   提醒逻辑','把概念数据转为便于理解的提示'),
            ('03','概念采集端','状态入口   /   照料记录','说明信息从哪里进入系统')]
    for i,(n,t,a,b) in enumerate(layers):
        y=139+i*105
        p.shape(43,y,651,89,'#FFFFFF',rounded=False)
        p.shape(43,y,174,89,GREEN if i==1 else '#DDEBDD',rounded=False)
        p.text(n,57,y+11,46,33,25,'#A5C5AD' if i==1 else '#739B7B',family=NUMERIC)
        p.text(t,57,y+47,147,33,18,'#FFFFFF' if i==1 else GREEN,bold=True)
        p.text(a,237,y+9,438,35,20,GREEN,bold=True)
        p.text(b,238,y+47,433,30,14.6,MUTED)
        if i<2:p.line(455,y+90,455,y+103,'#86AD8F',2)
    p.image('product-cutout.png',709,174,216,237,media)
    p.text('概念架构：信息组织与交互关系',45,463,863,28,13.6,MUTED)
    specs.append(p.finish())

    p=fresh(6,'#F8FBF6')
    p.text('把主体放大，让细节说话',42,35,866,68,34,GREEN,family=SERIF)
    p.line(228,203,490,258,'#86AD8F',1.3)
    p.line(647,315,730,230,'#86AD8F',1.3)
    p.line(466,397,467,442,'#86AD8F',1.3)
    p.image('product-cutout.png',248,99,470,354,media)
    p.text('01',45,143,89,50,38,'#7DA785',family=NUMERIC)
    p.text('状态入口',45,196,204,40,23,GREEN,bold=True)
    p.text('白色传感器概念造型\n用于说明采集入口',45,241,208,66,16,MUTED)
    p.text('02',737,144,120,50,38,CORAL,family=NUMERIC)
    p.text('手机端',737,198,181,40,23,GREEN,bold=True)
    p.text('以占位界面表达\n提醒与记录功能',737,244,181,65,16,MUTED)
    p.text('植物与花盆：养护对象',319,447,342,32,17,GREEN,center=True)
    specs.append(p.finish())

    p=fresh(7,GREEN,True)
    p.text('03',682,32,236,198,160,'#2D624E',family=NUMERIC)
    p.text('CARE EXPERIENCE',63,48,588,38,17,'#B9D8C2',family=NUMERIC)
    p.text('每一次照料',63,171,829,80,53,'#FFFFFF',family=PLAY)
    p.text('都值得被记住',63,259,829,80,53,'#FFB39E',family=PLAY)
    p.line(66,371,333,371,'#7FAD8A',1.3)
    p.text('用短句、留白与字体气质，完成章节转场。',66,391,816,45,19,'#D1E4D6')
    specs.append(p.finish())

    p=fresh(8,'#EFF5E9')
    p.text('接下来\n让体验更完整',43,59,259,116,34,GREEN,family=SERIF)
    p.text('后续设想',45,204,248,34,16,MUTED)
    p.image('product-cutout.png',21,283,260,180,media)
    p.line(319,135,319,457,'#92B79B',2)
    stages=[('01','先建立记录','让信息有清楚的查看入口。'),
            ('02','再优化提醒','探索更适合使用过程的提示。'),
            ('03','再完善体验','结合反馈继续调整交互细节。')]
    for i,(n,t,b) in enumerate(stages):
        x=354+i*50;y=127+i*116
        p.shape(312,y+21,14,14,CORAL if i==1 else '#8AB393',circle=True)
        p.line(329,y+28,x-12,y+28,'#92B79B',1)
        p.text(n,x,y,65,67,49,CORAL if i==1 else '#729B7C',family=NUMERIC)
        p.text(t,x+79,y+5,365,40,24,GREEN,bold=True)
        p.text(b,x+80,y+51,356,33,15.6,MUTED)
    specs.append(p.finish())

    sources={FONT:(assets/'demo-font.ttf',assets/'font-OFL.txt'),
             SERIF:(assets/'fonts/ZCOOLXiaoWei-Regular.ttf',assets/'fonts/zcoolxiaowei-OFL.txt'),
             PLAY:(assets/'fonts/ZCOOLKuaiLe-Regular.ttf',assets/'fonts/zcoolkuaile-OFL.txt'),
             NUMERIC:(assets/'fonts/BarlowCondensed-Regular.ttf',assets/'fonts/barlowcondensed-OFL.txt')}
    used={}
    for spec in specs:
        for e in spec['elements']:
            if e['elementType']=='text':
                c=e['content'];used[c['fontFamily']]=used.get(c['fontFamily'],'')+c['text']
    font_files={}
    for i,(family,text) in enumerate(used.items()):
        source,license_path=sources[family];f=TTFont(source)
        coverage=f.getBestCmap()
        missing=sorted({c for c in text if not c.isspace() and ord(c) not in coverage})
        if missing:raise ValueError(f'Font {family} is missing glyphs: {"".join(missing)}')
        options=subset.Options();options.name_IDs=['*'];options.name_legacy=True;options.name_languages=['*']
        worker=subset.Subsetter(options=options);worker.populate(text=text);worker.subset(f)
        relative=f'media/fonts/font-{i+1}.ttf';f.save(output/relative)
        shutil.copy2(license_path,fonts/f'font-{i+1}-OFL.txt');font_files[family]=relative
    brief={'project':'绿芽｜字体与版式选型图册','purpose':'比较排版表达，不要求正式稿混用全部方案',
           'design':{'mode':'explore','font_profiles':PROFILES,'constants':['绿色主题','清晰正文','概念素材']},
           'slides':[dict(id=i,title=row[0],layout=row[1],font_profile=row[2],density=row[3],
                          color_mode=row[4],focal_point=row[5],status='concept') for i,row in enumerate(PAGE_PLAN,1)]}
    (output/'design-plan.json').write_text(json.dumps(brief,ensure_ascii=False,indent=2),encoding='utf8')
    return write_project(prs,specs,output,'design-gallery',font_files,'绿芽｜字体与版式选型图册')


def main()->int:
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,required=True)
    parser.add_argument('--assets-dir',type=Path)
    parser.add_argument('--force',action='store_true')
    args=parser.parse_args()
    try:result=build(args.output_dir,args.assets_dir,args.force)
    except (ValueError,OSError) as exc:parser.error(str(exc))
    print(f'Built eight native layouts: {result}')
    print('Font embedding requires a supported presentation application; no fonts are installed by this script.')
    return 0


if __name__=='__main__':
    raise SystemExit(main())
